"""Разбор Excel-шаблона массовой загрузки: сборка флайтов, демо-строки, чистка чисел."""

from __future__ import annotations

from datetime import date, time
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.models.enums import EventStatus, GameType
from app.services.imports.bulk.assemble import MAX_LENGTHS, assemble_draft
from app.services.imports.bulk.reader import SheetFormatError, read_sheet
from app.services.imports.bulk.template import (
    COLUMNS,
    DEMO_ROW_FINGERPRINTS,
    row_fingerprint,
)
from tests.bulk_fixtures import TEMPLATE_PATH, build_workbook, default_rows, template_bytes


def parse(rows: list[dict[str, object]], **kwargs: object) -> object:
    return assemble_draft(read_sheet(build_workbook(rows, **kwargs)))  # type: ignore[arg-type]


def test_template_demo_fingerprints_match_shipped_file() -> None:
    """Константы отпечатков обязаны сходиться с docs/rasp_samples/day2_series_upload.xlsx."""
    read = read_sheet(template_bytes())
    actual = {row_fingerprint(row.values) for row in read.rows}
    assert actual == set(DEMO_ROW_FINGERPRINTS)


def test_downloadable_template_matches_the_reference_file() -> None:
    """Кнопка «Скачать шаблон» отдаёт ровно тот файл, под который написан парсер."""
    served = Path(__file__).resolve().parents[2] / "frontend" / "public" / TEMPLATE_PATH.name
    assert served.read_bytes() == TEMPLATE_PATH.read_bytes()


def test_shipped_template_imports_nothing() -> None:
    draft = assemble_draft(read_sheet(template_bytes()))
    assert draft.series == []
    assert draft.issues == []
    assert draft.demo_rows_skipped == [2, 3, 4, 5, 6]


def test_template_headers_cover_all_columns() -> None:
    read = read_sheet(template_bytes())
    assert read.unknown_headers == []
    assert read.rows
    known = {column.field for column in COLUMNS}
    assert set(read.rows[0].columns) == known


def test_template_instruction_lists_field_limits() -> None:
    """Лист «Инструкция» обязан повторять лимиты длины из загрузчика."""
    workbook = load_workbook(TEMPLATE_PATH, read_only=True, data_only=True)
    try:
        sheet = workbook["Инструкция"]
        described: dict[str, str] = {}
        for field_name, description in sheet.iter_rows(
            min_col=2, max_col=3, values_only=True
        ):
            if field_name in MAX_LENGTHS and isinstance(description, str):
                described[str(field_name)] = description
    finally:
        workbook.close()

    assert set(described) == set(MAX_LENGTHS)
    for field_name, limit in MAX_LENGTHS.items():
        assert str(limit) in described[field_name], field_name


def test_rows_group_into_series_events_and_flights() -> None:
    draft = parse(default_rows())

    assert [item.import_key for item in draft.series] == [
        "test-rpt-kaliningrad",
        "test-bpt-minsk",
    ]
    rpt, bpt = draft.series
    assert rpt.name == "Test RPT Kaliningrad"
    assert rpt.starts_on == date(2026, 8, 1)
    assert rpt.ends_on == date(2026, 8, 11)
    assert rpt.currency_code == "RUB"
    assert rpt.guarantee == Decimal("35000000.00")
    assert rpt.source_url == "https://t.me/rpt_poker/1234"

    assert [item.import_key for item in rpt.events] == ["4", "5"]
    knockout, main = rpt.events
    assert [flight.label for flight in knockout.flights] == ["Day 1A", "Day 1B"]
    assert [flight.play_time for flight in knockout.flights] == [time(12, 0), time(18, 0)]
    assert knockout.buyin == Decimal("14000.00")
    assert knockout.buyin_bounty == Decimal("6000.00")
    assert knockout.tags == ["bounty", "pko"]
    assert knockout.reentry_count == 2
    assert knockout.day_end_note == "till 12%"
    assert knockout.start_blinds == "100/200/200"
    assert knockout.game_type is GameType.NLH
    assert knockout.status is EventStatus.SCHEDULED

    assert [flight.label for flight in main.flights] == ["Day 1A", "Final Day"]
    assert [flight.level_minutes for flight in main.flights] == ["40", "45"]
    assert [flight.is_final for flight in main.flights] == [False, True]
    # Финальный день наследует бай-ин от строки со стартовым днём.
    assert main.buyin == Decimal("44000.00")
    assert main.notes == "Трансляция на канале серии"

    assert len(bpt.events) == 1
    single = bpt.events[0]
    assert [flight.label for flight in single.flights] == [None]
    assert single.flights[0].level_minutes == "25/20"


def test_empty_and_note_rows_are_skipped() -> None:
    rows = default_rows()
    draft = parse(rows)
    # Строка 7 пустая, строка 8 — сноска «↑ Строки 2–6 — ПРИМЕР…».
    assert draft.empty_rows_skipped == 2
    assert draft.rows_total == len(rows)


def test_demo_rows_are_skipped_but_edited_ones_are_kept() -> None:
    rows = default_rows()
    draft = assemble_draft(read_sheet(build_workbook(rows, keep_demo_rows=True)))

    assert draft.demo_rows_skipped == [2, 3, 4, 5, 6]
    assert {item.import_key for item in draft.series} == {
        "test-rpt-kaliningrad",
        "test-bpt-minsk",
    }


def test_edited_demo_row_is_imported_with_warning() -> None:
    """Кремовая заливка осталась, содержимое изменено — строка идёт в базу с оговоркой."""
    rows = default_rows()[:1]
    draft = assemble_draft(read_sheet(build_workbook(rows, keep_demo_rows=False)))
    assert len(draft.series) == 1

    workbook_rows = default_rows()[:1]
    workbook_rows[0]["series_key"] = "rpt-kaliningrad-2026-08"
    draft = parse(workbook_rows)
    codes = {issue.code for issue in draft.issues}
    assert "demo_key_reused" in codes
    assert len(draft.series) == 1


def test_money_with_currency_and_spaces_is_cleaned_with_warning() -> None:
    rows = default_rows()
    rows[2]["buyin"] = "44 000 ₽"
    draft = parse(rows)

    main = draft.series[0].events[1]
    assert main.buyin == Decimal("44000.00")
    warning = next(item for item in draft.issues if item.code == "money_cleaned")
    assert warning.severity == "warning"
    assert warning.row == 4
    assert warning.column == "T"


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("01.08.2026", date(2026, 8, 1)),
        ("2026-08-01", date(2026, 8, 1)),
        ("1/8/2026", date(2026, 8, 1)),
    ],
)
def test_date_formats(raw: str, expected: date) -> None:
    rows = default_rows()[:1]
    rows[0]["date"] = raw
    rows[0]["series_start"] = raw
    draft = parse(rows)
    assert draft.series[0].events[0].flights[0].play_date == expected


def test_unlimited_reentry() -> None:
    rows = default_rows()[:1]
    rows[0]["reentry"] = "безлимит"
    draft = parse(rows)
    event = draft.series[0].events[0]
    assert event.reentry_unlimited is True
    assert event.reentry_count is None


def test_series_key_is_case_and_space_insensitive() -> None:
    rows = default_rows()[:2]
    rows[1]["series_key"] = "  TEST-RPT-Kaliningrad "
    draft = parse(rows)
    assert len(draft.series) == 1
    assert draft.series[0].import_key == "test-rpt-kaliningrad"


def test_duplicate_flight_in_file_is_an_error() -> None:
    rows = default_rows()
    rows[1]["flight"] = "Day 1A"
    draft = parse(rows)

    issue = next(item for item in draft.issues if item.code == "duplicate_flight")
    assert issue.severity == "error"
    assert issue.row == 3
    assert issue.column == "Q"
    assert [flight.label for flight in draft.series[0].events[0].flights] == ["Day 1A"]


def test_series_field_mismatch_is_a_warning_and_first_row_wins() -> None:
    rows = default_rows()
    rows[1]["timezone"] = "Europe/Moscow"
    draft = parse(rows)

    issue = next(item for item in draft.issues if item.code == "row_mismatch")
    assert issue.severity == "warning"
    assert issue.row == 3
    assert issue.column == "G"
    assert draft.series[0].timezone == "Europe/Kaliningrad"


def test_missing_sheet_is_rejected() -> None:
    with pytest.raises(SheetFormatError) as exc:
        read_sheet(b"not an xlsx at all")
    assert exc.value.code == "unreadable_file"
