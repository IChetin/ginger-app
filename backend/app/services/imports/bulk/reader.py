"""Чтение листа «Турниры» в сырые строки с адресами ячеек."""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.imports.bulk.template import (
    COLUMN_BY_TITLE,
    DEMO_FILL_RGB,
    REQUIRED_FIELDS,
    SHEET_NAME,
)

HEADER_SEARCH_DEPTH = 10
# Значение длиннее ключа встречается только в сносках вроде «↑ Строки 2–6 — ПРИМЕР…».
NOTE_ROW_MIN_LENGTH = 64


class SheetFormatError(Exception):
    """Файл не является шаблоном Day2: нет листа или шапки."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class RawRow:
    row_no: int
    values: dict[str, Any]
    columns: dict[str, str]
    demo_fill: bool

    def get(self, field_name: str) -> Any:
        return self.values.get(field_name)

    def column(self, field_name: str) -> str | None:
        return self.columns.get(field_name)

    def text(self, field_name: str) -> str | None:
        value = self.values.get(field_name)
        if value is None:
            return None
        return " ".join(str(value).split()) or None


@dataclass
class SheetRead:
    rows: list[RawRow] = field(default_factory=list)
    unknown_headers: list[str] = field(default_factory=list)
    note_rows: list[int] = field(default_factory=list)
    empty_rows: int = 0


def _resolve_sheet(workbook: Any) -> Worksheet:
    if SHEET_NAME in workbook.sheetnames:
        return workbook[SHEET_NAME]
    raise SheetFormatError(
        "sheet_not_found",
        f"В файле нет листа «{SHEET_NAME}». Скачайте актуальный шаблон Day2.",
    )


def _find_header(sheet: Worksheet) -> tuple[int, dict[str, int], list[str]]:
    depth = min(HEADER_SEARCH_DEPTH, sheet.max_row or 1)
    for row_no in range(1, depth + 1):
        titles: dict[str, int] = {}
        unknown: list[str] = []
        for col_no in range(1, (sheet.max_column or 1) + 1):
            raw = sheet.cell(row_no, col_no).value
            if not isinstance(raw, str):
                continue
            title = " ".join(raw.split()).lower()
            column = COLUMN_BY_TITLE.get(title)
            if column is None:
                if title:
                    unknown.append(title)
                continue
            titles.setdefault(column.field, col_no)
        if "series_key" in titles and "event_key" in titles:
            return row_no, titles, unknown
    raise SheetFormatError(
        "header_not_found",
        "На листе «Турниры» не нашлась строка заголовков (нужны колонки series_key и event_key).",
    )


def read_sheet(data: bytes) -> SheetRead:
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl бросает разнородные ошибки
        raise SheetFormatError(
            "unreadable_file",
            "Файл не читается как XLSX. Сохраните шаблон в формате «Книга Excel (.xlsx)».",
        ) from exc

    try:
        sheet = _resolve_sheet(workbook)
        header_row, columns, unknown_headers = _find_header(sheet)

        missing = [item for item in REQUIRED_FIELDS if item not in columns]
        if missing:
            raise SheetFormatError(
                "missing_columns",
                "В шапке нет обязательных колонок: " + ", ".join(missing),
            )

        result = SheetRead(unknown_headers=unknown_headers)
        letters = {name: get_column_letter(index) for name, index in columns.items()}

        for row_no in range(header_row + 1, (sheet.max_row or header_row) + 1):
            values: dict[str, Any] = {}
            demo_fill = False
            for name, col_no in columns.items():
                cell = sheet.cell(row_no, col_no)
                raw = cell.value
                if isinstance(raw, str):
                    raw = raw.strip() or None
                if raw is None:
                    continue
                values[name] = raw
                fill = cell.fill
                if fill is not None and fill.patternType == "solid":
                    rgb = getattr(fill.fgColor, "rgb", None)
                    if isinstance(rgb, str) and rgb.upper() in DEMO_FILL_RGB:
                        demo_fill = True

            if not values:
                result.empty_rows += 1
                continue
            if len(values) == 1:
                only = next(iter(values.values()))
                if isinstance(only, str) and len(only) > NOTE_ROW_MIN_LENGTH:
                    result.note_rows.append(row_no)
                    continue

            # Буквы всех колонок шапки, а не только заполненных: адрес нужен и для
            # сообщения «обязательное поле не заполнено».
            result.rows.append(
                RawRow(
                    row_no=row_no,
                    values=values,
                    columns=letters,
                    demo_fill=demo_fill,
                )
            )
        return result
    finally:
        workbook.close()
